import base64
import io
from unittest.mock import patch

import openpyxl

from odoo.exceptions import UserError
from odoo.tests import TransactionCase, tagged


@tagged('post_install', '-at_install')
class TestProductImport(TransactionCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        if cls.env.ref('custom.comerciales', raise_if_not_found=False) is None:
            group = cls.env['res.groups'].create({'name': 'Comerciales'})
            cls.env['ir.model.data'].create({
                'module': 'custom', 'name': 'comerciales',
                'model': 'res.groups', 'res_id': group.id,
            })

    def _wizard(self, rows, prices=True, provider=False):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        headers = ['Marca', 'Referencia', 'Articulo', 'Descripcion',
                   'Familia', 'Subfamilia', 'Vendible']
        if prices:
            headers.append('Precio de venta')
        if provider:
            headers.append('Proveedor')
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        stream = io.BytesIO()
        workbook.save(stream)
        workbook.close()
        return self.env['product.import.wizard'].create({
            'excel_file': base64.b64encode(stream.getvalue()),
            'file_name': 'productos.xlsx',
        })

    def _row(self, reference, price=25):
        return ['Marca F6a', reference, 'Producto F6a', '', '', '', '00', price]

    def test_invalid_prices(self):
        for value in ('ilegible', 'NaN', 'inf', -1, None, '=1+1'):
            with self.subTest(value=value):
                wizard = self._wizard([self._row('F6A-INVALID', value)])
                with self.assertRaisesRegex(UserError, 'Fila 2, columna Precio de venta'):
                    wizard.action_import()

    def test_numeric_formats(self):
        wizard = self.env['product.import.wizard']
        for value, expected in [('1250.50', 1250.5), ('1.250,50', 1250.5),
                                ('1250,5', 1250.5), (1250.5, 1250.5)]:
            self.assertEqual(wizard._parse_price(value), expected)

    def test_missing_price_column_preserves_prices(self):
        product = self.env['product.template'].create({
            'name': 'Producto F6a', 'default_code': 'F6A-KEEP',
            'list_price': 123, 'standard_price': 45,
        })
        self._wizard([self._row(' f6a-keep ')[:-1]], prices=False).action_import()
        self.assertEqual(product.list_price, 123)
        self.assertEqual(product.standard_price, 45)

    def test_purchase_column_does_not_replace_sale_price(self):
        product = self.env['product.template'].create({
            'name': 'Producto F6a', 'default_code': 'F6A-COST', 'list_price': 123,
        })
        wizard = self._wizard([self._row('F6A-COST', 45)])
        workbook = openpyxl.load_workbook(io.BytesIO(base64.b64decode(wizard.excel_file)))
        workbook.active.cell(row=1, column=8, value='Precio de compra')
        stream = io.BytesIO()
        workbook.save(stream)
        workbook.close()
        wizard.excel_file = base64.b64encode(stream.getvalue())
        wizard.action_import()
        self.assertEqual(product.list_price, 123)
        self.assertEqual(product.standard_price, 45)

    def test_failure_after_one_hundred_templates_is_atomic(self):
        wizard = self._wizard([self._row('F6A-ATOMIC-%s' % index) for index in range(101)])
        model_class = type(wizard)
        original = model_class._get_or_create_categories
        calls = 0

        def fail_late(record, *args):
            nonlocal calls
            calls += 1
            if calls == 101:
                raise UserError('Fallo intermedio F6a')
            return original(record, *args)

        # Evitar que una regresión haga commit real sobre la base de tests.
        with patch.object(type(self.env.cr), 'commit') as commit:
            with patch.object(model_class, '_get_or_create_categories', fail_late):
                with self.assertRaisesRegex(UserError, 'Fallo intermedio F6a'):
                    wizard.action_import()
            commit.assert_not_called()
        self.assertFalse(self.env['product.template'].search([
            ('default_code', '=like', 'F6A-ATOMIC-%'),
        ]))

    def test_duplicate_reference(self):
        for reference in ('F6A-DUPLICATE', ' f6a-duplicate '):
            self.env['product.template'].create({'name': 'Duplicado', 'default_code': reference})
        with self.assertRaisesRegex(UserError, 'Referencia ambigua: F6A-DUPLICATE'):
            self._wizard([self._row('f6a-duplicate')]).action_import()

    def test_ambiguous_provider(self):
        for name in ('Proveedor F6a', ' PROVEEDOR F6A '):
            self.env['res.partner'].create({'name': name})
        with self.assertRaisesRegex(UserError, 'Proveedor ambiguo: PROVEEDOR F6A'):
            self._wizard([self._row('F6A-SUPPLIER') + [' proveedor f6a ']], provider=True).action_import()

    def test_reference_without_cached_value(self):
        for reference in (None, '=1+1'):
            with self.assertRaisesRegex(UserError, 'Fila 2, columna Referencia'):
                self._wizard([self._row(reference)]).action_import()

    def test_file_size_limit(self):
        wizard = self.env['product.import.wizard'].create({
            'excel_file': base64.b64encode(b'x' * (20 * 1024 * 1024 + 1)),
        })
        with self.assertRaisesRegex(UserError, '20 MB'):
            wizard.action_import()

    def test_row_limit(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.cell(row=50002, column=1, value='Exceso')
        with self.assertRaisesRegex(UserError, '50.000 filas'):
            self.env['product.import.wizard']._import_sheet(sheet)
        workbook.close()
