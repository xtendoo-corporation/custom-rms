# -*- coding: utf-8 -*-
import base64
import io
from unittest.mock import patch

import openpyxl

from odoo.exceptions import UserError
from odoo.tests.common import TransactionCase, tagged


@tagged('post_install', '-at_install')
class TestQuotationImport(TransactionCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        if cls.env.ref('custom.comerciales', raise_if_not_found=False) is None:
            group = cls.env['res.groups'].create({'name': 'Comerciales'})
            cls.env['ir.model.data'].create({
                'module': 'custom', 'name': 'comerciales',
                'model': 'res.groups', 'res_id': group.id,
            })
        cls.partner = cls.env['res.partner'].create({'name': 'Cliente prueba F6c'})
        cls.product = cls.env['product.product'].create({
            'name': 'Producto prueba F6c', 'default_code': 'F6C-PRODUCTO',
            'type': 'service', 'list_price': 10,
        })

    def _wizard(self, rows, header=None):
        # Generar el mismo fichero que recibe el asistente en producción.
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(header or ['Pedido', 'Cliente', 'Referencia producto', 'Cantidad', 'Precio'])
        for row in rows:
            sheet.append(row)
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        return self.env['rms.quotation.import.wizard'].create({
            'filename': 'prueba.xlsx', 'file': base64.b64encode(output.getvalue()),
        })

    def _row(self, ref='F6C-PEDIDO', price=10):
        return [ref, self.partner.name, self.product.default_code, 1, price]

    def test_reimport_rejected(self):
        wizard = self._wizard([self._row()])
        wizard.action_import()
        with self.assertRaisesRegex(UserError, 'ya existentes.*F6C-PEDIDO'):
            wizard.action_import()
        self.assertEqual(self.env['sale.order'].search_count([
            ('name', '=', 'F6C-PEDIDO'), ('company_id', '=', self.env.company.id)]), 1)

    def test_new_order_requires_customer(self):
        wizard = self._wizard([self._row(), ['F6C-OTRO', None, self.product.default_code, 1, 10]])
        with self.assertRaisesRegex(UserError, 'Fila 3.*cliente'):
            wizard.action_import()

    def test_localized_price(self):
        self._wizard([self._row(price='1.250,50')]).action_import()
        order = self.env['sale.order'].search([('name', '=', 'F6C-PEDIDO')])
        self.assertEqual(order.order_line.price_unit, 1250.5)

    def test_invalid_price(self):
        with self.assertRaisesRegex(UserError, 'Fila 2, columna Precio'):
            self._wizard([self._row(price='ilegible')]).action_import()

    def test_initial_reference_required(self):
        with self.assertRaisesRegex(UserError, 'Fila 2.*referencia'):
            self._wizard([self._row(ref=None)]).action_import()

    def test_ambiguous_header(self):
        with self.assertRaisesRegex(UserError, 'ambigua'):
            self._wizard([self._row()], ['Referencia', 'Cliente', 'Producto', 'Cantidad', 'Precio']).action_import()

    def test_invalid_numbers_and_ranges(self):
        for column, value in [(3, 'ilegible'), (3, 0), (3, -1), (5, 'ilegible'), (5, 101), (5, -1)]:
            row = self._row() + [0]
            row[column] = value
            with self.subTest(column=column, value=value), self.assertRaisesRegex(UserError, 'Fila 2, columna'):
                self._wizard([row], ['Pedido', 'Cliente', 'Referencia producto', 'Cantidad', 'Precio', 'Dto1']).action_import()

    def test_ambiguous_customer(self):
        self.env['res.partner'].create({'name': self.partner.name})
        with self.assertRaisesRegex(UserError, 'Fila 2.*cliente ambiguo'):
            self._wizard([self._row()]).action_import()

    def test_requested_state_required(self):
        for prefix in ('2M-', 'D-'):
            row = self._row()
            row[2] = prefix + self.product.default_code
            with self.subTest(prefix=prefix), self.assertRaisesRegex(UserError, 'Fila 2.*estado'):
                self._wizard([row]).action_import()

    def test_ambiguous_product(self):
        # copy() duplica la plantilla en Odoo 19; crear la variante explícitamente.
        duplicate = self.env['product.product'].create({
            'name': 'Producto duplicado F6c',
            'default_code': self.product.default_code,
            'type': 'service',
        })
        candidates = self.env['product.product'].search([
            ('default_code', '=', self.product.default_code),
        ])
        self.assertEqual(set(candidates.ids), {self.product.id, duplicate.id})
        self.assertFalse(candidates.product_template_attribute_value_ids)
        wizard = self._wizard([self._row()])
        # Verificar también la búsqueda directa, sin la precarga de action_import.
        with self.assertRaisesRegex(UserError, 'Fila 2.*ambigua'):
            wizard._find_product_variant(self.product.default_code, row_number=2)
        with self.assertRaisesRegex(UserError, 'Fila 2.*ambigua'):
            wizard.action_import()

    def test_file_size_limit(self):
        wizard = self._wizard([self._row()])
        wizard.file = base64.b64encode(b' ' * (20 * 1024 * 1024 + 1))
        with self.assertRaisesRegex(UserError, '20 MB'):
            wizard.action_import()

    def test_row_limit(self):
        wizard = self._wizard([self._row()])
        with patch.object(openpyxl, 'load_workbook') as load:
            load.return_value.active.iter_rows.return_value = iter(
                [('Pedido', 'Cliente')] + [('F6C', self.partner.name)] * 50001)
            with self.assertRaisesRegex(UserError, '50.000 filas'):
                wizard.action_import()
