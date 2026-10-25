import base64
import io
import json
from types import SimpleNamespace
from unittest.mock import patch

import openpyxl

from odoo.tests import tagged
from odoo.tests.common import HttpCase, TransactionCase, new_test_user

from odoo.addons.rms_global_equipment_map.controllers.main import GlobalEquipmentMapController


@tagged("post_install", "-at_install")
class TestMapRouteSecurity(HttpCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.portal = new_test_user(cls.env, login="s1_map_portal", password="s1-password", groups="base.group_portal")
        cls.internal = new_test_user(cls.env, login="s1_map_internal", password="s1-password", groups="base.group_user")
        cls.partner = cls.env["res.partner"].create({
            "name": "S1 cliente del mapa", "partner_latitude": 40.4,
            "partner_longitude": -3.7,
        })

    def _post(self, path, params):
        response = self.url_open(
            path,
            data=json.dumps({"jsonrpc": "2.0", "method": "call", "params": params, "id": 1}),
            headers={"Content-Type": "application/json"},
        )
        self.assertEqual(response.status_code, 200)
        return response.json()

    def test_portal_denied_before_read(self):
        self.authenticate(self.portal.login, "s1-password")
        with patch.object(GlobalEquipmentMapController, "_read_partners") as read:
            for path in ("partners", "export_xlsx"):
                with self.subTest(route=path):
                    result = self._post("/rms_global_equipment_map/" + path, {})
                    self.assertNotIn("result", result)
                    self.assertEqual(result["error"]["data"]["name"], "odoo.exceptions.AccessError")
            read.assert_not_called()

    def test_internal_can_read_and_export(self):
        self.authenticate(self.internal.login, "s1-password")
        result = self._post("/rms_global_equipment_map/partners", {})
        self.assertNotIn("error", result)
        self.assertIn(self.partner.id, [row["id"] for row in result["result"]])
        result = self._post("/rms_global_equipment_map/export_xlsx", {"partner_ids": [self.partner.id]})
        self.assertNotIn("error", result)
        workbook = openpyxl.load_workbook(io.BytesIO(base64.b64decode(result["result"]["content"])))
        self.assertEqual(workbook.active.cell(2, 1).value, self.partner.name)


@tagged("post_install", "-at_install")
class TestMapCompanySecurity(TransactionCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.company_a = cls.env["res.company"].create({"name": "F3 compañía A"})
        cls.company_b = cls.env["res.company"].create({"name": "F3 compañía B"})
        cls.user = new_test_user(
            cls.env, login="f3_global_company_a", groups="base.group_user",
            company_id=cls.company_a.id, company_ids=[(6, 0, cls.company_a.ids)],
        )
        cls.partners = cls.env["res.partner"].create([
            {"name": "F3 cliente A", "company_id": cls.company_a.id,
             "partner_latitude": 40.4, "partner_longitude": -3.7},
            {"name": "F3 cliente B", "company_id": cls.company_b.id,
             "partner_latitude": 41.4, "partner_longitude": 2.1},
        ])

    def _request(self):
        return patch(
            "odoo.addons.rms_global_equipment_map.controllers.main.request",
            SimpleNamespace(env=self.env["res.partner"].with_user(self.user).with_context(
                allowed_company_ids=self.company_a.ids,
            ).env),
        )

    def test_map_respects_company_rules(self):
        with self._request():
            rows = GlobalEquipmentMapController().get_global_equipment_map_partners()
        ids = {row["id"] for row in rows}
        self.assertIn(self.partners[0].id, ids)
        self.assertNotIn(self.partners[1].id, ids)

    def test_export_filters_arbitrary_partner_ids(self):
        with self._request():
            result = GlobalEquipmentMapController().export_global_equipment_map_xlsx(
                self.partners.ids,
            )
        workbook = openpyxl.load_workbook(io.BytesIO(base64.b64decode(result["content"])))
        names = [row[0] for row in workbook.active.iter_rows(min_row=2, values_only=True)]
        self.assertEqual(names, [self.partners[0].name])

    def test_map_hides_restricted_child_contact_details(self):
        contact = self.env["res.partner"].create({
            "name": "F3 contacto privado", "parent_id": self.partners[0].id,
            "email": "privado@example.com", "phone": "123456789",
        })
        self.env["ir.rule"].create({
            "name": "F3 ocultar contacto hijo", "model_id": self.env.ref("base.model_res_partner").id,
            "domain_force": "[('id', '!=', %s)]" % contact.id,
        })
        with self._request():
            rows = GlobalEquipmentMapController().get_global_equipment_map_partners()
        row = next(row for row in rows if row["id"] == self.partners[0].id)
        self.assertEqual(row["contact_name"], "")
        self.assertEqual(row["email"], "")
        self.assertEqual(row["phone"], "")
